"""Estrazione del testo dai formati supportati, a partire dal contenuto binario."""

from __future__ import annotations

import csv
import io
import logging
from collections.abc import Callable
from dataclasses import dataclass

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class Section:
    """Blocco di testo estratto con la sua posizione logica nel documento."""

    text: str
    location: str = ""


class ExtractionError(Exception):
    """Errore durante l'estrazione del testo da un documento."""


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_plain_text(data: bytes) -> list[Section]:
    return [Section(_decode_text(data))]


def extract_csv(data: bytes) -> list[Section]:
    text = _decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(io.StringIO(text), dialect)
    lines = [" | ".join(cell.strip() for cell in row if cell.strip()) for row in rows]
    return [Section("\n".join(line for line in lines if line))]


def extract_pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # noqa: BLE001 - pypdf solleva eccezioni eterogenee
            raise ExtractionError("PDF protetto da password") from exc
    sections = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - una pagina corrotta non deve bloccare il resto
            logger.debug("Pagina %d illeggibile", number, exc_info=True)
            continue
        if text.strip():
            sections.append(Section(text, f"pagina {number}"))
    return sections


def extract_docx(data: bytes) -> list[Section]:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                # Una riga di tabella diventa una frase autonoma
                parts.append(" | ".join(dict.fromkeys(cells)) + ".")
    return [Section("\n\n".join(parts))]


def extract_xlsx(data: bytes) -> list[Section]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections = []
    try:
        for sheet in workbook.worksheets:
            lines = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if cells:
                    lines.append(" | ".join(cells) + ".")
            if lines:
                sections.append(Section("\n".join(lines), f"foglio {sheet.title}"))
    finally:
        workbook.close()
    return sections


def extract_pptx(data: bytes) -> list[Section]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    sections = []
    for number, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        texts.append(" | ".join(cells) + ".")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                texts.append(notes)
        if texts:
            sections.append(Section("\n\n".join(texts), f"slide {number}"))
    return sections


Extractor = Callable[[bytes], list[Section]]

_TEXT_EXTENSIONS = (
    ".txt",
    ".md",
    ".log",
    ".ini",
    ".cfg",
    ".conf",
    ".json",
    ".xml",
    ".yaml",
    ".yml",
    ".rst",
)


def _build_registry() -> dict[str, Extractor]:
    registry: dict[str, Extractor] = {ext: extract_plain_text for ext in _TEXT_EXTENSIONS}
    registry[".csv"] = extract_csv
    # I formati binari sono registrati solo se la libreria è installata
    optional = {
        ".pdf": ("pypdf", extract_pdf),
        ".docx": ("docx", extract_docx),
        ".xlsx": ("openpyxl", extract_xlsx),
        ".xlsm": ("openpyxl", extract_xlsx),
        ".pptx": ("pptx", extract_pptx),
    }
    for ext, (module, extractor) in optional.items():
        try:
            __import__(module)
        except ImportError:
            logger.warning("Libreria '%s' assente: formato %s non supportato", module, ext)
            continue
        registry[ext] = extractor
    return registry


EXTRACTORS: dict[str, Extractor] = _build_registry()


def supported_extensions() -> set[str]:
    return set(EXTRACTORS)


def extract_sections(data: bytes, extension: str) -> list[Section]:
    """Estrae le sezioni di testo; solleva ExtractionError in caso di problemi."""
    extractor = EXTRACTORS.get(extension.lower())
    if extractor is None:
        raise ExtractionError(f"Formato non supportato: {extension}")
    try:
        return [s for s in extractor(data) if s.text.strip()]
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - file corrotti generano eccezioni di ogni tipo
        raise ExtractionError(f"{type(exc).__name__}: {exc}") from exc
