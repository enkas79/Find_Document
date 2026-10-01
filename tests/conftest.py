"""Fixture condivise: cartella di documenti tecnici di esempio in vari formati."""

from __future__ import annotations

import os
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


@pytest.fixture
def docs_folder(tmp_path: Path) -> Path:
    folder = tmp_path / "docs"
    (folder / "manuali").mkdir(parents=True)

    (folder / "pompa.txt").write_text(
        "Manuale della pompa centrifuga P-200.\n\n"
        "La pressione massima di esercizio della pompa è 16 bar. "
        "La temperatura del fluido non deve superare 90 °C.\n\n"
        "Manutenzione: sostituire le guarnizioni ogni 2000 ore di funzionamento.",
        encoding="utf-8",
    )
    (folder / "manuali" / "flangia.md").write_text(
        "# Montaggio flangia\n\n"
        "Le viti M12 della flangia vanno serrate con una coppia di serraggio di 85 Nm "
        "seguendo uno schema a croce.\n\n"
        "Le viti M16 richiedono 210 Nm.",
        encoding="utf-8",
    )

    import docx

    document = docx.Document()
    document.add_paragraph("Specifiche del motore elettrico.")
    document.add_paragraph("Il motore deve essere alimentato a 400 V trifase con protezione IP55.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Potenza nominale"
    table.rows[0].cells[1].text = "7,5 kW"
    document.save(folder / "motore.docx")

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Valvole"
    sheet.append(["Codice", "Descrizione", "Diametro"])
    sheet.append(["V-101", "Valvola a sfera in acciaio inox", "DN50"])
    workbook.save(folder / "valvole.xlsx")

    (folder / "~$temporaneo.docx").write_bytes(b"lock")  # file di lock Office da ignorare
    (folder / "corrotto.pdf").write_bytes(b"%PDF-1.4 non valido")
    (folder / "immagine.png").write_bytes(b"\x89PNG")  # formato non supportato
    return folder
